import hashlib
import io
import math
import re
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from zoneinfo import ZoneInfo
from openpyxl import load_workbook
from sqlalchemy import delete, select
from .database import Asset, Snapshot, MarketCache
from .rules import evaluate, norm

STOCK_HEADERS=['Ticker','Empresa','Setor','Tipo de ação','Tag Along','Free Float','Segmento de listagem','Governo como Majoritário?','Patrim Liq','Liquidez média diária','Margem EBIT','Margem Líquida','ROIC','ROE','Anos desde o IPO','Anos gerando lucro','Dívida Líquida/EBITDA']
FII_HEADERS=['Cód do fundo','Tipo de fundo','Segmento de Atuação','Liquidez diária','Patrimônio Líquido','Quantidade de Imóveis','Multi-inquilino','Vacância','% em CRIs','Quantidade de CRIs','Administrador','Tempo de listagem','Tipo de Gestão','Taxa de adm','Taxa de Performance']
STOCK_NUMBERS=set(STOCK_HEADERS)-{'Ticker','Empresa','Setor','Tipo de ação','Segmento de listagem','Governo como Majoritário?'}
FII_NUMBERS={'Liquidez diária','Patrimônio Líquido','Quantidade de Imóveis','Vacância','% em CRIs','Quantidade de CRIs','Tempo de listagem','Taxa de adm'}
MONTHS=['janeiro','fevereiro','marco','abril','maio','junho','julho','agosto','setembro','outubro','novembro','dezembro']


class ImportErrorDetail(ValueError):
    pass


def number(v,label,row):
    if v is None or (isinstance(v,str) and v.strip() in ('','-','—')):return None
    if isinstance(v,bool):raise ImportErrorDetail(f'Linha {row}: {label} inválido.')
    if isinstance(v,str):
        v=v.strip();pct=v.endswith('%');v=v.removesuffix('%').replace('R$','').strip()
        if ',' in v:v=v.replace('.','').replace(',','.')
        try:v=float(v)/(100 if pct else 1)
        except ValueError:raise ImportErrorDetail(f'Linha {row}: {label} inválido.')
    if not isinstance(v,(float,int)) or not math.isfinite(v):raise ImportErrorDetail(f'Linha {row}: {label} inválido.')
    return v


def parse_workbook(content,kind,reference,filename):
    if kind not in ('stocks','fiis'):raise ImportErrorDetail('Tipo de planilha inválido.')
    if not re.fullmatch(r'20\d\d-(0[1-9]|1[0-2])',reference):raise ImportErrorDetail('Competência deve ser AAAA-MM.')
    now=datetime.now(ZoneInfo('America/Sao_Paulo')).strftime('%Y-%m')
    if reference>now:raise ImportErrorDetail('A competência não pode ser futura.')
    fn=norm(filename)
    found=re.search(r'('+'|'.join(MONTHS)+r')[\s_-]*(20\d{2})',fn)
    if found:
        inferred=f'{found[2]}-{MONTHS.index(found[1])+1:02}'
        if reference!=inferred:raise ImportErrorDetail(f'O nome do arquivo indica {inferred}, diferente de {reference}.')
    if not filename.lower().endswith('.xlsx'):raise ImportErrorDetail('Envie um arquivo .xlsx.')
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            infos=z.infolist()
            if len(infos)>2500 or sum(i.file_size for i in infos)>60*1024*1024:raise ImportErrorDetail('Planilha descompactada excede o limite de segurança.')
            if any('vbaproject' in i.filename.lower() for i in infos):raise ImportErrorDetail('Arquivos com macros não são aceitos.')
        book=load_workbook(io.BytesIO(content),read_only=True,data_only=True,keep_links=False)
    except ImportErrorDetail:raise
    except Exception:raise ImportErrorDetail('Arquivo XLSX inválido ou corrompido.')
    required=STOCK_HEADERS if kind=='stocks' else FII_HEADERS
    numeric=STOCK_NUMBERS if kind=='stocks' else FII_NUMBERS
    records=[];seen=set()
    try:
        sheet=book['Método Davi'] if 'Método Davi' in book.sheetnames else book.worksheets[0]
        header=None;indexes=None
        for row_index,values in enumerate(sheet.iter_rows(max_row=6000,max_col=70,values_only=True),1):
            if header is None:
                if row_index>15:break
                normalized={norm(v):i for i,v in enumerate(values) if isinstance(v,str) and v.strip()}
                if all(norm(k) in normalized for k in required):
                    header=values;indexes={k:normalized[norm(k)] for k in required}
                    for k in ['Div. Yield','CAGR do lucro em 5 anos','DY 12M Acumulado']:
                        if norm(k) in normalized:indexes[k]=normalized[norm(k)]
                continue
            ticker=str(values[indexes[required[0]]] or '').strip().upper()
            if not ticker:continue
            if not re.fullmatch(r'[A-Z][A-Z0-9]{3}\d{1,2}',ticker):raise ImportErrorDetail(f'Linha {row_index}: ticker inválido {ticker[:20]}.')
            if ticker in seen:raise ImportErrorDetail(f'Ticker repetido: {ticker}. A importação inteira foi cancelada.')
            seen.add(ticker)
            d={k:values[i] for k,i in indexes.items()}
            d[required[0]]=ticker
            for k in numeric:d[k]=number(d[k],k,row_index)
            for k in set(required)-numeric:
                d[k]=str(d[k] or '').strip()
            for k in ('Tag Along','Free Float','Vacância','% em CRIs','Taxa de adm'):
                if k in d and d[k] is not None and not 0<=d[k]<=1:raise ImportErrorDetail(f'Linha {row_index}: {k} deve estar entre 0 e 1 (ou texto com %).')
            for k in ('Anos gerando lucro','Anos desde o IPO','Quantidade de Imóveis','Quantidade de CRIs','Tempo de listagem'):
                if k in d and d[k] is not None and (d[k]<0 or d[k]!=int(d[k])):raise ImportErrorDetail(f'Linha {row_index}: {k} deve ser inteiro não negativo.')
            for k in ('Div. Yield','CAGR do lucro em 5 anos','DY 12M Acumulado'):
                if k in d:d[k]=number(d[k],k,row_index)
            d['_reference']=reference
            if kind=='stocks':
                name=d['Empresa'];category='stock';segment=d['Setor']
            else:
                name=ticker;typ=norm(d['Tipo de fundo']);category={'papel':'paper','tijolo':'brick','hibrido':'hybrid'}.get(typ,'other');segment=d['Segmento de Atuação']
            records.append(dict(ticker=ticker,kind=kind,name=name,category=category,segment=segment,source_row=row_index,data=d,assessment=evaluate(kind,d)))
            if len(records)>4000:raise ImportErrorDetail('Limite de 4 mil ativos por planilha.')
        if header is None:raise ImportErrorDetail('Cabeçalhos não encontrados. Use a Planilha de Ouro de '+('ações' if kind=='stocks' else 'FIIs')+'.')
        if not records:raise ImportErrorDetail('A planilha não contém ativos válidos.')
        if sheet.max_row>6000:raise ImportErrorDetail('Planilha com mais de 6 mil linhas; remova linhas excedentes.')
    finally:book.close()
    return records


def replace_snapshot(db,kind,reference,filename,content,records,user_id,expected_revision):
    snap=db.execute(select(Snapshot).where(Snapshot.kind==kind).with_for_update()).scalar_one()
    if snap.revision!=expected_revision:raise ImportErrorDetail('A base mudou durante a revisão. Atualize e confira a importação novamente.')
    if snap.reference and reference<snap.reference:raise ImportErrorDetail('Não é permitido substituir a base atual por uma competência anterior.')
    digest=hashlib.sha256(content).hexdigest()
    if snap.sha256==digest and snap.reference==reference:return snap,False
    db.execute(delete(Asset).where(Asset.kind==kind))
    for r in records:db.add(Asset(**r))
    snap.reference=reference;snap.filename=Path(filename).name;snap.sha256=digest
    snap.imported_at=datetime.now(timezone.utc);snap.imported_by=user_id
    snap.count=len(records);snap.revision+=1
    db.flush()
    # Cache de preços não é a planilha e pode ser mantido; apenas remove tickers órfãos.
    db.execute(delete(MarketCache).where(~MarketCache.ticker.in_(select(Asset.ticker))))
    return snap,True
